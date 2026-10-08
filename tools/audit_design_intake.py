"""Read-only attachment audit. Never edits the source workbook or game saves."""
from pathlib import Path
import collections
import hashlib
import json
import openpyxl

root = Path(__file__).resolve().parents[1]
downloads = Path('C:/Users/godsi/Downloads')
source = downloads/'Guildborne_UXUI_Asset_Registry.xlsx'
workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
result = {'source': str(source), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'sheets': []}
for sheet in workbook:
    rows = [[v for v in row] for row in sheet.iter_rows(values_only=True)]
    populated = [(i+1,row) for i,row in enumerate(rows) if any(v is not None for v in row)]
    result['sheets'].append({'name':sheet.title,'rows':rows})
    print(json.dumps({'sheet':sheet.title, 'rows':sheet.max_row, 'columns':sheet.max_column,
                      'firstRows':populated[:3]},ensure_ascii=True,default=str))
out = root/'docs/uat01/intake'
out.mkdir(parents=True,exist_ok=True)
(out/'registry-source.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
for source in (downloads/'Guildborne_Progression_Quest_Design_v2/data').glob('*.json'):
    data=json.loads(source.read_text(encoding='utf-8-sig'))
    (out/source.name).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'file':source.name,'structure':{k:len(v) if isinstance(v,(dict,list)) else v for k,v in data.items()}},ensure_ascii=True))
workbook.close()
